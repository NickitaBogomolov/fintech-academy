import io
from datetime import date
from decimal import Decimal

import pytest
from docx import Document

from dfa_studio.documents import formatting
from dfa_studio.documents.docx_export import markdown_to_docx
from dfa_studio.documents.render import (
    Operator,
    TemplateStore,
    TemplateValidationError,
    render_document,
    validate_template,
)
from dfa_studio.domain.schedule import build_schedule

OPERATOR = Operator(name="ООО «Оператор ИС»", ogrn="0000000000010", website="https://operator.example")


def test_number_formatting() -> None:
    assert formatting.money(Decimal("15.29")) == "15,29"
    assert formatting.money_words(Decimal(1000)) == "1 000 (одна тысяча) российских рублей"
    assert formatting.int_words(300_000) == "300 000 (триста тысяч)"
    assert formatting.percent_words(Decimal("0.18")) == "18,00% (восемнадцать процентов)"
    assert formatting.date_words(date(2026, 12, 15)) == "15 декабря 2026 г."


@pytest.fixture
def store(tmp_path) -> TemplateStore:
    return TemplateStore(tmp_path)


def test_builtin_template_renders_spec_values(store, spec, calendar) -> None:
    source = store.get("decision_coupon").source
    text = render_document(source, spec, build_schedule(spec, calendar), OPERATOR, date(2026, 12, 1))
    assert "Выпуск № **SVR-001-2026**" in text
    assert "300 000 (триста тысяч) штук" in text
    assert "18,00% (восемнадцать процентов) годовых" in text
    assert "| 5 | 15.04.2027 | 15.05.2027 | 30 | 14,79 | — | 17.05.2027 * |" in text
    assert "**89,76**" in text
    assert "только лицами, являющимися квалифицированными инвесторами" in text


def test_docx_contains_tables_and_headings(store, spec, calendar) -> None:
    text = render_document(
        store.get("decision_coupon").source, spec, build_schedule(spec, calendar), OPERATOR, date(2026, 12, 1)
    )
    document = Document(io.BytesIO(markdown_to_docx(text)))
    assert document.paragraphs[0].text == "РЕШЕНИЕ О ВЫПУСКЕ ЦИФРОВЫХ ФИНАНСОВЫХ АКТИВОВ"
    assert len(document.tables) == 2
    assert document.tables[1].cell(1, 4).text == "15,29"


def test_unknown_variables_are_rejected() -> None:
    with pytest.raises(TemplateValidationError, match="client"):
        validate_template("{{ client.name }}")


def test_user_template_lifecycle(store) -> None:
    info = store.save("short-form", "{#- title: Краткая форма -#}\nНоминал {{ spec.nominal | money }}")
    assert info.title == "Краткая форма"
    assert info.variables == ["spec.nominal"]
    assert any(item.id == "short-form" for item in store.list())
    store.delete("short-form")
    assert all(item.id != "short-form" for item in store.list())


def test_builtin_template_cannot_be_overwritten(store) -> None:
    with pytest.raises(TemplateValidationError):
        store.save("decision_coupon", "x")


def test_percent_noun_agrees_with_number() -> None:
    assert formatting.percent_words(Decimal("0.21")).endswith("(двадцать один процент)")
    assert formatting.percent_words(Decimal("0.22")).endswith("(двадцать два процента)")
    assert formatting.percent_words(Decimal("0.185")).endswith("(восемнадцать целых пять десятых процента)")
