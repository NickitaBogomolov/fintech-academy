import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph

BOLD = re.compile(r"\*\*(.+?)\*\*")
TABLE_DIVIDER = re.compile(r"^\|\s*:?-{3,}")


def _add_runs(paragraph: Paragraph, text: str) -> None:
    text = text.replace("\\*", "*")
    position = 0
    for match in BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _flush_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            paragraph = table.cell(row_index, col_index).paragraphs[0]
            _add_runs(paragraph, value)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def markdown_to_docx(markdown: str) -> bytes:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    table_rows: list[list[str]] = []
    first_heading = True
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("|"):
            if not TABLE_DIVIDER.match(line):
                table_rows.append(_cells(line))
            continue
        _flush_table(document, table_rows)
        table_rows = []
        if not line.strip():
            continue
        if line.startswith("# "):
            heading = document.add_heading(line[2:], level=0 if first_heading else 1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first_heading = False
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            _add_runs(document.add_paragraph(style="List Bullet"), line[2:])
        else:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(paragraph, line)
    _flush_table(document, table_rows)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
